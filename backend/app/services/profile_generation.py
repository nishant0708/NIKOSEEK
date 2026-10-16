from pathlib import Path
import re

from fastapi import HTTPException, status
from pypdf import PdfReader
from docx import Document

from app.services.llm import generate_text
from app.services.storage import download_file, resume_belongs_to, save_details_markdown, save_profile_markdown

URL_PATTERN = re.compile(r"(?i)(?<![@\w])(?:https?://|www\.)?[a-z0-9](?:[a-z0-9-]*[a-z0-9])?(?:\.[a-z0-9](?:[a-z0-9-]*[a-z0-9])?)+(?:/[^\s<>\"']*)?")
URL_TRAILING_PUNCTUATION = ".,;:!?)]}"


def _extract_resume_text(filename: str, data: bytes) -> str:
    try:
        if filename.endswith(".pdf"):
            from io import BytesIO
            return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
        if filename.endswith(".docx"):
            from io import BytesIO
            return "\n".join(paragraph.text for paragraph in Document(BytesIO(data)).paragraphs)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Could not read text from this resume") from exc
    raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Only PDF and DOCX resumes are supported")


def _extract_resume_links(filename: str, data: bytes, text: str = "") -> list[str]:
    """Collect visible URLs plus links embedded as PDF annotations or DOCX hyperlinks."""
    found: list[str] = []
    for match in URL_PATTERN.findall(text):
        url = match.rstrip(URL_TRAILING_PUNCTUATION)
        if url and not url.lower().startswith(("http://", "https://")):
            url = "https://" + url
        if url:
            found.append(url)
    try:
        from io import BytesIO
        if filename.endswith(".pdf"):
            for page in PdfReader(BytesIO(data)).pages:
                for annotation in page.get("/Annots") or []:
                    action = annotation.get_object().get("/A")
                    uri = action.get("/URI") if action else None
                    if uri and str(uri).lower().startswith(("http://", "https://")):
                        found.append(str(uri).strip())
        elif filename.endswith(".docx"):
            document = Document(BytesIO(data))
            for relation in document.part.rels.values():
                if relation.reltype.endswith("/hyperlink") and relation.target_ref:
                    target = relation.target_ref.strip()
                    if target.lower().startswith(("http://", "https://")):
                        found.append(target)
    except Exception:
        # Link extraction is best-effort; normal resume text extraction owns validation.
        pass
    return list(dict.fromkeys(link.rstrip(URL_TRAILING_PUNCTUATION) for link in found if link))


def generate_profile_documents(username: str, resume_key: str, projects: str, achievements: str) -> tuple[str, str]:
    if not resume_belongs_to(username, resume_key):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="That resume does not belong to you")
    resume_bytes = download_file(resume_key)
    extension = Path(resume_key).suffix.lower()
    resume_text = _extract_resume_text(extension, resume_bytes)
    resume_links = _extract_resume_links(extension, resume_bytes, resume_text)
    links_context = "\n".join(f"- {link}" for link in resume_links) or "(No links found)"
    source = f"RESUME:\n{resume_text[:30000]}\n\nLINKS FOUND IN RESUME:\n{links_context}\n\nUSER PROJECT DETAILS:\n{projects}\n\nUSER ACHIEVEMENTS:\n{achievements}"
    details = generate_text(source, "Create details.md for a job seeker. Preserve facts only. Use Markdown headings for projects, achievements, skills, experience evidence and measurable impact. Do not invent information.")
    profile = generate_text(source + "\n\nGENERATED DETAILS:\n" + details, "Create a concise, editable profile.md for job matching. Include headline, summary, skills, experience, projects, education and job preferences. Preserve facts only; do not invent information.")
    if resume_links:
        links_section = "\n\n## Links found in resume\n" + "\n".join(f"- {link}" for link in resume_links)
        details += links_section
        profile += links_section
    save_details_markdown(username, details)
    save_profile_markdown(username, profile)
    return details, profile
